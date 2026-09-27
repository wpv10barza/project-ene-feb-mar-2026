import os
import sys
import time
import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from docx import Document
from docx.shared import Inches
from io import BytesIO

# --- LÓGICA DE ARGUMENTOS ---
def obtener_ruta():
    ruta_base = r"D:\118102025\Downloads"
    nombre_base = "terminos gemini.docx"
    if len(sys.argv) < 2: return os.path.join(ruta_base, nombre_base)
    arg = sys.argv[1]
    if "\\" in arg or "/" in arg:
        return arg if arg.lower().endswith(".docx") else os.path.join(arg, nombre_base)
    return os.path.join(ruta_base, arg if arg.lower().endswith(".docx") else arg + ".docx")

RUTA_WORD = obtener_ruta()

def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}")

def iniciar_driver():
    options = uc.ChromeOptions()
    perfil = os.path.join(os.getcwd(), "uc_profile_final")
    options.add_argument(f"--user-data-dir={perfil}")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--no-sandbox")
    # FORZAMOS LA VERSIÓN 145 PARA EVITAR EL CONFLICTO
    return uc.Chrome(options=options, version_main=145)

try:
    log(f"Iniciando (v145)... Destino: {RUTA_WORD}")
    if not os.path.exists(os.path.dirname(RUTA_WORD)): os.makedirs(os.path.dirname(RUTA_WORD))
    
    driver = iniciar_driver()
    doc = Document()
    driver.get("https://gemini.google.com")

    log("Carga el chat y presiona ENTER en esta terminal...")
    input(">>> [ENTER] PARA PROCESAR")

    ventana_gemini = driver.current_window_handle
    nodos = driver.find_elements(By.CSS_SELECTOR, "user-query, model-response")
    empezar = False
    
    for nodo in nodos:
        if "Actúa como un experto" in nodo.text: empezar = True
        if not empezar: continue

        # Extraer Texto
        for el in nodo.find_elements(By.CSS_SELECTOR, "p, li, pre"):
            txt = el.text.strip()
            if txt: doc.add_paragraph(txt)

        # Buscar enlaces de interés
        links = [l.get_attribute("href") for l in nodo.find_elements(By.TAG_NAME, "a") if l.get_attribute("href") and ("github.com" in l.get_attribute("href") or "researchgate.net" in l.get_attribute("href"))]
        
        for url in links:
            try:
                log(f"Accediendo a: {url}")
                driver.switch_to.new_window('tab')
                driver.get(url)
                time.sleep(5)
                
                target = None
                # Selectores mejorados para ResearchGate y GitHub
                for sel in ["img[src*='figure']", "article img", "img.nova-legacy-e-image__graphic", ".repository-content img"]:
                    imgs = driver.find_elements(By.CSS_SELECTOR, sel)
                    for im in imgs:
                        if im.size['width'] > 150:
                            target = im
                            break
                    if target: break
                
                if target:
                    doc.add_picture(BytesIO(target.screenshot_as_png), width=Inches(5.5))
                    doc.add_paragraph(f"Fuente: {url}")
                    log("   [OK] Captura guardada.")
                
                driver.close()
                driver.switch_to.window(ventana_gemini)
            except Exception as e:
                log(f"   [!] Error en pestaña: {e}")
                if len(driver.window_handles) > 1: driver.close()
                driver.switch_to.window(ventana_gemini)

    doc.save(RUTA_WORD)
    log(f"TODO LISTO. Archivo en: {RUTA_WORD}")
except Exception as e:
    log(f"ERROR CRÍTICO: {e}")
finally:
    if 'driver' in locals(): driver.quit()
